# 第一个“真活”：写一个 Excel，再读回来，并算出平均分
# 运行前先装库：python -m pip install openpyxl
from pathlib import Path

from openpyxl import Workbook, load_workbook

out = Path(__file__).parent / "成绩单.xlsx"

# 1. 造数据并写出
wb = Workbook()
ws = wb.active
ws.title = "成绩"
ws.append(["姓名", "语文", "数学", "英语"])

rows = [
    ("深田蓝", 88, 92, 79),
    ("隙月", 76, 95, 88),
    ("流小莹", 91, 68, 84),
]
for row in rows:
    ws.append(row)
wb.save(out)
print(f"已写入：{out}")

# 2. 读回来，逐行打印，并算每个人的平均分
wb2 = load_workbook(out)
ws2 = wb2["成绩"]
print("\n读回来的内容：")
for row in ws2.iter_rows(min_row=2, values_only=True):
    name, chinese, math, english = row
    avg = (chinese + math + english) / 3
    print(f"{name}：语文{chinese} 数学{math} 英语{english}，平均 {avg:.1f} 分")

print("\n搞定：文件的写和读，是所有 Excel 自动化的起点。")
